"""Synthetic DOCX grids through the existing public Document contract."""

import hashlib
import os

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from fileforge.document import compare_versions, read_document
from fileforge.extractor import docx


def read_saved(document, path):
    document.save(path)
    before = (path.read_bytes(), path.stat().st_mtime_ns)
    result = read_document(str(path))
    assert before == (path.read_bytes(), path.stat().st_mtime_ns)
    return result


def test_mixed_block_order_and_full_table_content(tmp_path):
    source = tmp_path / "中文 空格.DOCX"
    d = Document()
    d.add_paragraph("before")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "middle"
    d.add_paragraph("after")
    text = read_saved(d, source)["data"]["document"]["text"]
    assert (
        text.index("before")
        < text.index("[TABLE]")
        < text.index("middle")
        < text.index("[/TABLE]")
        < text.index("after")
    )


def test_table_only_success_and_identity_unchanged(tmp_path):
    source = tmp_path / "only.docx"
    d = Document()
    table = d.add_table(rows=2, cols=3)
    for row, values in zip(
        table.rows,
        [("Category", "Level", "Points"), ("Example", "Gold", "7")],
        strict=True,
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    result = read_saved(d, source)
    assert result["success"], result
    document = result["data"]["document"]
    assert 'R2: C1="Example" | C2="Gold" | C3="7"' in document["text"]
    assert document["metadata"] == {}
    expected = (
        "file_"
        + hashlib.sha256(os.path.normcase(str(source.resolve())).encode()).hexdigest()
    )
    assert document["file_id"] == expected
    assert read_document(str(source)) == result
    assert set(document) == {
        "file_id",
        "path",
        "filename",
        "extension",
        "text",
        "size",
        "modified_time",
        "metadata",
    }


@pytest.mark.parametrize("height,width", [(1, 2), (2, 1), (2, 2)])
def test_merged_anchor_once_and_covered_cells_reference_it(tmp_path, height, width):
    d = Document()
    table = d.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(height - 1, width - 1)).text = "UNIQUE_CATEGORY"
    table.cell(2, 0).text = "Second category"
    table.cell(2, 1).text = "Silver"
    table.cell(2, 2).text = "4"
    text = read_saved(d, tmp_path / "merged.docx")["data"]["document"]["text"]
    assert text.count("UNIQUE_CATEGORY") == 1
    assert f'C1(rowspan={height},colspan={width})="UNIQUE_CATEGORY"' in text
    assert text.count("@R1C1") == height * width - 1
    assert 'R3: C1="Second category" | C2="Silver" | C3="4"' in text


def test_vertical_merges_keep_category_level_and_points_aligned(tmp_path):
    d = Document()
    t = d.add_table(rows=4, cols=4)
    for c, v in zip(
        t.rows[0].cells, ["Category", "Level", "Rank", "Points"], strict=True
    ):
        c.text = v
    t.cell(1, 0).merge(t.cell(3, 0)).text = "Activity"
    t.cell(1, 1).merge(t.cell(2, 1)).text = "Local"
    t.cell(3, 1).text = "Regional"
    for r, rank, score in [(1, "First", "3"), (2, "Second", "1.5"), (3, "First", "8")]:
        t.cell(r, 2).text, t.cell(r, 3).text = rank, score
    text = read_saved(d, tmp_path / "scores.docx")["data"]["document"]["text"]
    assert 'R3: C1=@R2C1 | C2=@R2C2 | C3="Second" | C4="1.5"' in text
    assert 'R4: C1=@R2C1 | C2="Regional" | C3="First" | C4="8"' in text
    assert text.count("Activity") == 1


def test_empty_cells_internal_spaces_and_delimiters_are_unambiguous(tmp_path):
    d = Document()
    t = d.add_table(rows=1, cols=3)
    t.cell(0, 0).text = "  A  B | C\nD  "
    t.cell(0, 2).text = "5"
    text = read_saved(d, tmp_path / "space.docx")["data"]["document"]["text"]
    assert 'C1="A  B | C\\nD" | C2="" | C3="5"' in text


def test_nested_table_stays_inside_cell_in_block_order(tmp_path):
    d = Document()
    outer = d.add_table(rows=1, cols=2)
    cell = outer.cell(0, 0)
    cell.text = "before nested"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "inner"
    cell.add_paragraph("after nested")
    outer.cell(0, 1).text = "neighbor"
    text = read_saved(d, tmp_path / "nested.docx")["data"]["document"]["text"]
    assert text.count("[TABLE]") == 2 and text.count("[/TABLE]") == 2
    assert (
        text.index("before nested")
        < text.index("inner")
        < text.index("after nested")
        < text.index("neighbor")
    )
    assert 'C2="neighbor"' in text


@pytest.mark.parametrize("with_paragraph", [False, True])
def test_whitespace_only_tables_do_not_turn_empty_document_into_success(
    tmp_path, with_paragraph
):
    d = Document()
    if with_paragraph:
        d.add_paragraph(" \t ")
    t = d.add_table(rows=2, cols=2)
    for row in t.rows:
        for cell in row.cells:
            cell.text = "  \t "
    result = read_saved(d, tmp_path / "empty.docx")
    assert not result["success"] and result["data"] is None
    assert "body paragraphs or tables" in result["error"]


def test_long_table_tail_and_source_preserved_and_handles_released(tmp_path):
    source = tmp_path / "long.docx"
    d = Document()
    d.add_table(rows=1, cols=1).cell(0, 0).text = "x" * 3000 + "END_OF_TABLE"
    result = read_saved(d, source)
    assert result["success"] and "END_OF_TABLE" in result["data"]["document"]["text"]
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    assert read_document(str(source)) == result
    moved = source.with_name("temporary-name.docx")
    source.rename(moved)
    moved.rename(source)
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)
    assert len(docx.extract(source)) == 2000  # upstream snippet API remains bounded


def test_added_table_text_counts_against_existing_version_limit(tmp_path):
    d = Document()
    d.add_paragraph("body")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "x" * 10001 + "TAIL"
    result = read_saved(d, tmp_path / "oversize.docx")
    assert result["success"]
    document = result["data"]["document"]
    assert "TAIL" in document["text"] and len(document["text"]) > 10000
    compared = compare_versions([document])
    assert not compared["success"] and compared["data"] is None
    assert "10000" in compared["error"]


def test_omitted_grid_cells_not_shifted_left(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=3)
    row = t.rows[1]
    row._tr.remove(row._tr.tc_lst[0])
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    row._tr.get_or_add_trPr().append(before)
    from docx.table import _Cell

    _Cell(row._tr.tc_lst[0], t).text = "Level"
    _Cell(row._tr.tc_lst[1], t).text = "9"
    text = read_saved(d, tmp_path / "omitted.docx")["data"]["document"]["text"]
    assert 'R2: C1=[omitted] | C2="Level" | C3="9"' in text


@pytest.mark.parametrize("leading_row", [False, True])
def test_invalid_vertical_continuation_is_standard_failure(tmp_path, leading_row):
    d = Document()
    t = d.add_table(rows=2 if leading_row else 1, cols=1)
    t.cell(0, 0).text = "not a merge origin"
    merge = OxmlElement("w:vMerge")
    t.cell(1 if leading_row else 0, 0)._tc.get_or_add_tcPr().append(merge)
    result = read_saved(d, tmp_path / "invalid.docx")
    assert result["success"] is False and result["data"] is None
    assert "parsing failed" in result["error"]


def test_blank_table_does_not_create_fake_content_beside_paragraph(tmp_path):
    d = Document()
    d.add_paragraph("body")
    d.add_table(rows=2, cols=2)
    assert (
        read_saved(d, tmp_path / "blank-table.docx")["data"]["document"]["text"]
        == "body"
    )


def test_legacy_horizontal_merge_fails_instead_of_inventing_grid(tmp_path):
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "legacy"
    merge = OxmlElement("w:hMerge")
    merge.set(qn("w:val"), "restart")
    t.cell(0, 0)._tc.get_or_add_tcPr().append(merge)
    result = read_saved(d, tmp_path / "legacy.docx")
    assert result["success"] is False and result["data"] is None
    assert "Legacy DOCX hMerge" in result["error"]
