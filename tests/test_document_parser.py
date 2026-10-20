"""Document contract tests using disposable real source files."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

import pytest

from fileforge.document import read_document


@pytest.mark.parametrize(
    ("content", "encoding"),
    [
        ("报名材料\n" * 1000 + "全文结尾", "utf-8"),
        ("Résumé professionnel\n" * 20, "latin-1"),
        ("中文报名材料\n" * 20, "utf-16"),
    ],
)
def test_full_text_contract_and_source_unchanged(
    tmp_dir: Path, content: str, encoding: str
) -> None:
    source = tmp_dir / "HARNESS_document.TXT"
    raw = content.encode(encoding)
    source.write_bytes(raw)
    before = source.stat()
    result = read_document(str(source))
    assert set(result) == {"success", "tool", "data", "message", "error"}
    assert result["success"] is True
    assert result["tool"] == "read_document"
    assert result["error"] is None
    document = result["data"]["document"]
    assert set(document) == {
        "file_id",
        "path",
        "filename",
        "extension",
        "text",
        "size",
        "modified_time",
        "metadata",
    }
    assert (
        document["file_id"]
        == "file_"
        + hashlib.sha256(
            os.path.normcase(str(source.resolve())).encode("utf-8")
        ).hexdigest()
    )
    assert document["path"] == str(source.resolve())
    assert document["filename"] == source.name
    assert document["extension"] == ".txt"
    assert document["text"] == content
    assert document["size"] == len(raw)
    assert datetime.fromisoformat(document["modified_time"]) == datetime.fromtimestamp(
        before.st_mtime, tz=UTC
    )
    assert isinstance(document["metadata"], dict)
    json.dumps(result)
    assert source.read_bytes() == raw
    assert source.stat().st_mtime_ns == before.st_mtime_ns


@pytest.mark.parametrize("content", ["", " \n\t "])
def test_empty_content_is_explicit_failure(tmp_dir: Path, content: str) -> None:
    source = tmp_dir / "HARNESS_empty.txt"
    source.write_text(content)
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "empty" in result["error"].lower()
    assert source.read_text() == content


@pytest.mark.parametrize("extension", [".xlsm", ".md", ".bin"])
def test_unsupported_formats_are_rejected(tmp_dir: Path, extension: str) -> None:
    source = tmp_dir / ("HARNESS_other" + extension)
    source.write_bytes(b"sample")
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "unsupported" in result["error"].lower()
    assert source.read_bytes() == b"sample"


def test_missing_file(tmp_dir: Path) -> None:
    result = read_document(str(tmp_dir / "HARNESS_missing.txt"))
    assert result["success"] is False
    assert result["data"] is None
    assert "not found" in result["error"].lower()


def test_directory_is_not_a_document(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_directory.txt"
    source.mkdir()
    result = read_document(str(source))
    assert result["success"] is False
    assert "regular file" in result["error"].lower()


@pytest.mark.parametrize("error", [PermissionError("denied"), OSError("I/O failure")])
def test_read_failure(tmp_dir: Path, error: OSError) -> None:
    source = tmp_dir / "HARNESS_unreadable.txt"
    source.write_text("source remains intact")
    with patch("fileforge.document.parser.text.extract", side_effect=error):
        result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert result["error"]
    assert source.read_text() == "source remains intact"


def test_expanduser_uses_existing_text_reader(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_home.txt"
    source.write_text("expanded path content", encoding="utf-8")
    before = source.stat().st_mtime_ns
    with patch("os.path.expanduser", return_value=str(tmp_dir)):
        result = read_document("~/HARNESS_home.txt")
    assert result["success"] is True
    assert result["data"]["document"]["file_id"] == (
        read_document(str(source))["data"]["document"]["file_id"]
    )
    assert result["data"]["document"]["path"] == str(source.resolve())
    assert result["data"]["document"]["text"] == "expanded path content"
    assert source.read_text(encoding="utf-8") == "expanded path content"
    assert source.stat().st_mtime_ns == before


@pytest.mark.parametrize(
    "error", [OSError("cannot resolve"), RuntimeError("link loop")]
)
def test_path_resolution_failure(error: Exception) -> None:
    with patch.object(Path, "resolve", side_effect=error):
        result = read_document("HARNESS_unresolved.txt")
    assert result["success"] is False
    assert result["data"] is None
    assert "read failed" in result["error"].lower()


def test_identity_survives_repeat_reads_and_content_changes(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_identity.txt"
    source.write_text("original", encoding="utf-8")
    first = read_document(str(source))["data"]["document"]
    repeated = read_document(str(source))["data"]["document"]
    assert first["file_id"] == repeated["file_id"]
    source.write_text("updated text", encoding="utf-8")
    updated = read_document(str(source))["data"]["document"]
    assert first["file_id"] == updated["file_id"]
    assert updated["text"] == "updated text"


def test_same_length_change_with_restored_mtime_keeps_identity_not_content(
    tmp_dir: Path,
) -> None:
    """Path identity and unchanged stat fields do not certify unchanged bytes."""
    source = tmp_dir / "HARNESS_same_stat.txt"
    original = b"alpha evidence\n"
    replacement = b"bravo evidence\n"
    assert original != replacement and len(original) == len(replacement)
    source.write_bytes(original)
    first_result = read_document(str(source))
    assert first_result["success"], first_result
    assert read_document(str(source)) == first_result
    first = first_result["data"]["document"]
    before = source.stat()
    original_digest = hashlib.sha256(source.read_bytes()).hexdigest()

    source.write_bytes(replacement)
    os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns))
    second_result = read_document(str(source))
    assert second_result["success"], second_result
    assert read_document(str(source)) == second_result
    second = second_result["data"]["document"]

    for field in ("file_id", "path", "size", "modified_time"):
        assert first[field] == second[field], field
    assert source.stat().st_mtime_ns == before.st_mtime_ns
    assert source.stat().st_size == before.st_size == len(replacement)
    assert first["text"] == original.decode("utf-8")
    assert second["text"] == replacement.decode("utf-8")
    assert first["text"] != second["text"]
    assert source.read_bytes() == replacement
    assert original_digest != hashlib.sha256(source.read_bytes()).hexdigest()


def test_copy_and_rename_change_identity(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_source.txt"
    source.write_text("same content", encoding="utf-8")
    original = read_document(str(source))["data"]["document"]
    copied = tmp_dir / "HARNESS_copy.txt"
    shutil.copyfile(source, copied)
    copy_doc = read_document(str(copied))["data"]["document"]
    assert copy_doc["text"] == original["text"]
    assert copy_doc["file_id"] != original["file_id"]
    renamed = tmp_dir / "HARNESS_renamed.txt"
    source.rename(renamed)
    renamed_doc = read_document(str(renamed))["data"]["document"]
    assert renamed_doc["file_id"] != original["file_id"]
    destination = tmp_dir / "HARNESS_subdirectory"
    destination.mkdir()
    moved = destination / renamed.name
    renamed.rename(moved)
    moved_doc = read_document(str(moved))["data"]["document"]
    assert moved_doc["file_id"] != renamed_doc["file_id"]
    assert moved_doc["text"] == original["text"]


def test_relative_absolute_and_symlink_paths_agree(
    tmp_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_dir / "HARNESS_target.txt"
    source.write_text("target content", encoding="utf-8")
    monkeypatch.chdir(tmp_dir)
    absolute = read_document(str(source))["data"]["document"]
    relative = read_document("./HARNESS_target.txt")["data"]["document"]
    assert absolute == relative
    link = tmp_dir / "HARNESS_alias.txt"
    try:
        link.symlink_to(source)
    except (OSError, NotImplementedError):
        pytest.skip("Symlinks are not available in this environment")
    assert read_document(str(link))["data"]["document"] == absolute


@pytest.mark.skipif(os.name == "nt", reason="Linux case-sensitive path semantics")
def test_linux_does_not_lowercase_paths(tmp_dir: Path) -> None:
    upper = tmp_dir / "HARNESS_CASE.txt"
    lower = tmp_dir / "HARNESS_case.txt"
    upper.write_text("same", encoding="utf-8")
    lower.write_text("same", encoding="utf-8")
    assert read_document(str(upper))["data"]["document"]["file_id"] != (
        read_document(str(lower))["data"]["document"]["file_id"]
    )


def test_utf8_path_encoding_failure_returns_envelope() -> None:
    with patch.object(Path, "resolve", return_value=Path("/HARNESS_\udcff.txt")):
        result = read_document("HARNESS_bad_path.txt")
    assert result["success"] is False
    assert result["data"] is None
    assert result["tool"] == "read_document"
    assert result["error"]


def test_broken_symlink_returns_failure(tmp_dir: Path) -> None:
    link = tmp_dir / "HARNESS_broken.txt"
    try:
        link.symlink_to(tmp_dir / "HARNESS_missing.txt")
    except (OSError, NotImplementedError):
        pytest.skip("Symlinks are not available in this environment")
    result = read_document(str(link))
    assert result["success"] is False
    assert result["data"] is None
    assert "not found" in result["error"].lower()


def _make_pdf(path: Path, pages: list[list[str]]) -> None:
    """Write a real, uncompressed PDF with Helvetica text and an xref table.

    ASCII fixtures deliberately avoid font installation and extra PDF writers.
    An empty page has no text layer, without asserting it is a scanned page.
    """
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            "<< /Type /Pages /Kids ["
            + " ".join(f"{4 + i * 2} 0 R" for i in range(len(pages)))
            + f"] /Count {len(pages)} >>"
        ).encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, lines in enumerate(pages):
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                "/Resources << /Font << /F1 3 0 R >> >> "
                f"/Contents {5 + index * 2} 0 R >>"
            ).encode()
        )
        commands = ["BT /F1 10 Tf 40 750 Td 12 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj T*")
        commands.append("ET")
        stream = "\n".join(commands).encode("ascii")
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    raw = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(raw)
    raw.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        raw.extend(f"{offset:010d} 00000 n \n".encode())
    raw.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode()
    )
    path.write_bytes(raw)


def test_pdf_full_text_multiple_pages_and_source_preserved(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_long.pdf"
    _make_pdf(
        source, [["FIRST_PAGE " + "abc " * 15] * 40, ["SECOND_PAGE", "PDF_END_MARKER"]]
    )
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    result = read_document(str(source))
    assert result["success"] is True, result
    document = result["data"]["document"]
    assert len(document["text"]) > 2000
    assert "FIRST_PAGE" in document["text"]
    assert document["text"].endswith("PDF_END_MARKER")
    assert document["extension"] == ".pdf"
    assert document == read_document(str(source))["data"]["document"]
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)


def test_pdf_no_text_is_not_diagnosed_as_scanned(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_blank.pdf"
    _make_pdf(source, [[]])
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "No extractable text" in result["error"]
    assert "does not determine" in result["error"]


def test_pdf_corrupt_is_parse_failure(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_corrupt.pdf"
    source.write_bytes(b"not a PDF")
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "failed" in result["error"]
    assert "No extractable text" not in result["error"]
    assert source.read_bytes() == b"not a PDF"


def _make_docx(path: Path, paragraphs: list[str], table: bool = False) -> None:
    """Create a real Office document using the existing runtime dependency."""
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_ONLY_MARKER"
    document.save(path)


def test_docx_long_body_paragraphs_and_source_preserved(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_long.docx"
    _make_docx(source, ["正文第一段" * 1000, "第二段", "DOCX_END_MARKER"], table=True)
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    result = read_document(str(source))
    assert result["success"] is True, result
    document = result["data"]["document"]
    assert document["text"].startswith("正文第一段")
    assert "第二段\nDOCX_END_MARKER\n[TABLE]" in document["text"]
    assert len(document["text"]) > 2000
    assert 'C1="TABLE_ONLY_MARKER"' in document["text"]
    assert document["extension"] == ".docx"
    assert document == read_document(str(source))["data"]["document"]
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)


@pytest.mark.parametrize("table", [False, True])
def test_docx_no_body_text_and_optional_table(tmp_dir: Path, table: bool) -> None:
    source = tmp_dir / "HARNESS_no_body.docx"
    _make_docx(source, ["   "], table=table)
    result = read_document(str(source))
    if table:
        assert result["success"] is True
        assert 'C1="TABLE_ONLY_MARKER"' in result["data"]["document"]["text"]
    else:
        assert result["success"] is False
        assert result["data"] is None
        assert "No extractable text" in result["error"]
        assert "body paragraphs or tables" in result["error"]


def test_docx_corrupt_is_parse_failure(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_corrupt.docx"
    source.write_bytes(b"not a DOCX")
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "failed" in result["error"]
    assert "No extractable text" not in result["error"]


def _make_xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    """Generate a real workbook with the existing openpyxl dependency."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_xlsx_full_text_multiple_sheets_and_source_preserved(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_long.xlsx"
    _make_xlsx(
        source, {"预算": [["long cell " * 1000]], "Final": [["XLSX_END_MARKER", 123]]}
    )
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    result = read_document(str(source))
    assert result["success"] is True, result
    document = result["data"]["document"]
    assert len(document["text"]) > 2000
    assert "[预算]" in document["text"]
    assert "[Final]" in document["text"]
    assert document["text"].endswith("XLSX_END_MARKER | 123")
    assert document["extension"] == ".xlsx"
    assert document == read_document(str(source))["data"]["document"]
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)


@pytest.mark.parametrize("rows", [[], [["=1+2"]]])
def test_xlsx_no_values_or_uncached_formula_is_explicit_failure(
    tmp_dir: Path, rows: list[list]
) -> None:
    source = tmp_dir / "HARNESS_no_values.xlsx"
    _make_xlsx(source, {"Empty": rows})
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "No extractable cell values" in result["error"]
    assert "cached" in result["error"]
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)


def test_xlsx_formula_cache_is_used_without_calculation(tmp_dir: Path) -> None:
    from xml.etree import ElementTree
    from zipfile import ZipFile

    source = tmp_dir / "HARNESS_formula.xlsx"
    _make_xlsx(source, {"Formula": [["=1+2", "=5+6", "visible label"]]})
    # Fixture setup only: insert a deliberately stale cache (99, not 3).
    # The reader must preserve that cached value and never compute 3 or 11.
    with ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    sheet = ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    sheet.find(".//s:c[@r='A1']/s:v", namespace).text = "99"
    entries["xl/worksheets/sheet1.xml"] = ElementTree.tostring(sheet)
    with ZipFile(source, "w") as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    result = read_document(str(source))
    assert result["success"] is True, result
    assert result["data"]["document"]["text"] == "[Formula]\n99 | visible label"
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)


def test_xlsx_corrupt_is_parse_failure(tmp_dir: Path) -> None:
    source = tmp_dir / "HARNESS_corrupt.xlsx"
    source.write_bytes(b"not a workbook")
    result = read_document(str(source))
    assert result["success"] is False
    assert result["data"] is None
    assert "failed" in result["error"]
    assert "No extractable" not in result["error"]


def _write_document(path: Path, content: str) -> None:
    if path.suffix == ".pdf":
        _make_pdf(path, [[content]])
    elif path.suffix == ".docx":
        _make_docx(path, [content])
    elif path.suffix == ".xlsx":
        _make_xlsx(path, {"Content": [[content]]})
    else:
        path.write_text(content, encoding="utf-8")


@pytest.mark.parametrize("extension", [".txt", ".pdf", ".docx", ".xlsx"])
def test_all_formats_share_contract_identity_and_read_only_behavior(
    tmp_dir: Path, monkeypatch: pytest.MonkeyPatch, extension: str
) -> None:
    source = tmp_dir / ("HARNESS_identity" + extension)
    _write_document(source, "original marker")
    monkeypatch.chdir(tmp_dir)
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    first = read_document(str(source))
    assert first["success"] is True, first
    assert set(first) == {"success", "tool", "data", "message", "error"}
    assert first["tool"] == "read_document"
    assert first["error"] is None
    document = first["data"]["document"]
    assert set(document) == {
        "file_id",
        "path",
        "filename",
        "extension",
        "text",
        "size",
        "modified_time",
        "metadata",
    }
    assert document["path"] == str(source.resolve())
    assert document["filename"] == source.name
    assert document["extension"] == extension
    assert document["size"] == len(before[0])
    assert datetime.fromisoformat(document["modified_time"]) == datetime.fromtimestamp(
        source.stat().st_mtime, tz=UTC
    )
    assert document["metadata"] == {}
    expected_id = (
        "file_"
        + hashlib.sha256(
            os.path.normcase(str(source.resolve())).encode("utf-8")
        ).hexdigest()
    )
    assert document["file_id"] == expected_id
    assert read_document(source.name) == first
    json.dumps(first)
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)
    copied = tmp_dir / ("HARNESS_copy" + extension)
    shutil.copyfile(source, copied)
    assert read_document(str(copied))["data"]["document"]["file_id"] != expected_id
    _write_document(source, "updated marker")
    before_update_read = (source.read_bytes(), source.stat().st_mtime_ns)
    updated = read_document(str(source))["data"]["document"]
    assert updated["file_id"] == expected_id
    assert "updated marker" in updated["text"]
    assert before_update_read == (source.read_bytes(), source.stat().st_mtime_ns)


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".xlsx"])
def test_format_read_permission_failure_is_not_success(
    tmp_dir: Path, extension: str
) -> None:
    source = tmp_dir / ("HARNESS_denied" + extension)
    _write_document(source, "unchanged")
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    module = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx"}[extension]
    with patch(
        f"fileforge.document.parser.{module}.extract",
        side_effect=PermissionError("denied"),
    ):
        result = read_document(str(source))
    assert result == {
        "success": False,
        "tool": "read_document",
        "data": None,
        "message": "Document read failed",
        "error": "Permission denied reading document",
    }
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)


def test_docx_unextracted_property_changes_bytes_not_full_document(
    tmp_dir: Path,
) -> None:
    """A parsed snapshot is not a digest of all DOCX archive properties."""
    import copy
    from io import BytesIO
    from zipfile import ZIP_STORED, ZipFile

    from docx import Document

    from fileforge.document import compare_versions

    source = tmp_dir / "HARNESS_properties.docx"
    document = Document()
    document.add_paragraph("Unchanged extracted body")
    document.core_properties.title = "alpha"
    document.save(source)
    template = source.read_bytes()

    def archive_with_title(title: bytes) -> bytes:
        # Store the property entry without compression, ensuring equal-length
        # title bytes cannot accidentally change the public Document.size.
        buffer = BytesIO()
        with ZipFile(BytesIO(template)) as original, ZipFile(buffer, "w") as output:
            for info in original.infolist():
                info = copy.copy(info)
                data = original.read(info.filename)
                if info.filename == "docProps/core.xml":
                    assert data.count(b"<dc:title>alpha</dc:title>") == 1
                    data = data.replace(
                        b"<dc:title>alpha</dc:title>",
                        b"<dc:title>" + title + b"</dc:title>",
                    )
                    info.compress_type = ZIP_STORED
                output.writestr(info, data)
        return buffer.getvalue()

    original = archive_with_title(b"alpha")
    replacement = archive_with_title(b"bravo")
    assert original != replacement and len(original) == len(replacement)
    source.write_bytes(original)
    assert Document(source).core_properties.title == "alpha"
    first = read_document(str(source))
    assert first["success"], first
    before = source.stat()
    source.write_bytes(replacement)
    os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns))
    assert Document(source).core_properties.title == "bravo"
    second = read_document(str(source))
    assert second["success"], second
    # Establish the requested real-file fixture condition before comparison.
    assert first["data"]["document"] == second["data"]["document"]
    assert first["data"]["document"]["metadata"] == {}
    assert source.stat().st_mtime_ns == before.st_mtime_ns
    assert (
        hashlib.sha256(original).digest()
        != hashlib.sha256(source.read_bytes()).digest()
    )
    preserved = (source.read_bytes(), source.stat().st_mtime_ns)
    compared = compare_versions([first["data"]["document"], second["data"]["document"]])
    assert compared["success"] and compared["error"] is None
    group = compared["data"]["groups"][0]
    assert group["documents"] == [first["data"]["document"]]
    assert group["newest_candidate"] is None
    assert group["recommendation_reason"] == "single_document"
    assert preserved == (source.read_bytes(), source.stat().st_mtime_ns)
