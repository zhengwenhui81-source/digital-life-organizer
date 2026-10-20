"""C-owned checks of delivered evidence helpers, not Package execution."""

import hashlib
import os
from copy import copy, deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile

import pytest
from docx import Document

from fileforge.agent.package_evidence import (
    bind_selected_snapshots,
    revalidate_documents,
)
from fileforge.document import read_document
from fileforge.package import plan_package


@pytest.fixture
def source_read(tmp_path):
    source = tmp_path / "material.txt"
    source.write_text("Temporary evidence", encoding="utf-8")
    result = read_document(str(source))
    assert result["success"], result
    return source, result


@pytest.mark.parametrize(
    "left,right",
    [
        ({"nested": {"items": [1, None]}}, {"nested": {"items": [1]}}),
        ({"nested": {}}, {"nested": {"missing": None}}),
        ({"nested": [True]}, {"nested": [1]}),
        ({"nested": [1]}, {"nested": [1.0]}),
        ({"nested": [1, 2]}, {"nested": [2, 1]}),
    ],
)
def test_nested_snapshot_differences_reject_without_mutation(
    tmp_path, source_read, left, right
):
    _, original = source_read
    first, second = deepcopy(original), deepcopy(original)
    first["data"]["document"]["metadata"] = left
    second["data"]["document"]["metadata"] = right
    inputs = [(0, first), (1, second)]
    before = deepcopy(inputs)
    with pytest.raises(ValueError, match="Contradictory"):
        revalidate_documents(str(tmp_path), inputs)
    assert inputs == before


def test_nested_key_order_matches_complete_reread(tmp_path, source_read):
    _, original = source_read
    retained, current = deepcopy(original), deepcopy(original)
    retained["data"]["document"]["metadata"] = {"a": {"x": 1, "y": None}, "b": [True]}
    current["data"]["document"]["metadata"] = {"b": [True], "a": {"y": None, "x": 1}}
    before = deepcopy((retained, current))
    # Synthetic metadata tests comparison only; the real reader emits {}.
    checked = revalidate_documents(
        str(tmp_path), [(4, retained)], reader=lambda **_: deepcopy(current)
    )
    assert set(checked) == {4}
    assert (retained, current) == before


def test_docx_equal_parse_byte_change_during_real_reread_is_rejected(tmp_path):
    source = tmp_path / "properties.docx"
    doc = Document()
    doc.add_paragraph("Identical extracted body")
    doc.core_properties.title = "alpha"
    doc.save(source)
    template = source.read_bytes()

    def archive(title):
        buffer = BytesIO()
        with ZipFile(BytesIO(template)) as original, ZipFile(buffer, "w") as output:
            for entry in original.infolist():
                entry = copy(entry)
                data = original.read(entry.filename)
                if entry.filename == "docProps/core.xml":
                    assert data.count(b"<dc:title>alpha</dc:title>") == 1
                    data = data.replace(b"<dc:title>alpha</dc:title>", title)
                    entry.compress_type = ZIP_STORED
                output.writestr(entry, data)
        return buffer.getvalue()

    first = archive(b"<dc:title>alpha</dc:title>")
    second = archive(b"<dc:title>bravo</dc:title>")
    assert len(first) == len(second)
    assert hashlib.sha256(first).digest() != hashlib.sha256(second).digest()
    source.write_bytes(first)
    original = read_document(str(source))
    assert original["success"], original
    before = source.stat()

    def changing_reader(**kwargs):
        result = read_document(**kwargs)
        source.write_bytes(second)
        os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns))
        assert Document(source).core_properties.title == "bravo"
        assert read_document(str(source)) == result == original
        assert source.stat().st_size == before.st_size
        assert source.stat().st_mtime_ns == before.st_mtime_ns
        return result

    assert Document(source).core_properties.title == "alpha"
    with pytest.raises(ValueError, match="source bytes changed during reread"):
        revalidate_documents(str(tmp_path), [(0, original)], reader=changing_reader)


@pytest.mark.parametrize("indices", [[-1], [True], ["0"], [0.0], [0, 0]])
def test_invalid_or_duplicate_indices_reject(tmp_path, source_read, indices):
    _, original = source_read
    with pytest.raises(ValueError, match="index is invalid"):
        revalidate_documents(str(tmp_path), [(index, original) for index in indices])


@pytest.mark.parametrize("path_kind", ["relative", "outside"])
def test_path_outside_workspace_or_relative_reject(tmp_path, source_read, path_kind):
    source, original = source_read
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    altered = deepcopy(original)
    altered["data"]["document"]["path"] = (
        source.name if path_kind == "relative" else str(source)
    )
    with pytest.raises(ValueError, match="escapes workspace"):
        revalidate_documents(str(workspace), [(0, altered)])


@pytest.mark.parametrize(
    "section,field,value",
    [
        ("source", "path", "other.txt"),
        ("source", "filename", "other.txt"),
        ("snapshot", "size", -1),
        ("snapshot", "mtime_ns", -1),
        ("snapshot", "size", True),
        ("snapshot", "mtime_ns", None),
    ],
)
def test_selected_snapshot_field_mismatch_rejects(
    tmp_path, source_read, section, field, value
):
    _, original = source_read
    checked = revalidate_documents(str(tmp_path), [(3, original)])
    evidence = checked[3]
    # Helper input only: this is not a D planner result or executable preview.
    plan = {
        "selected_files": [
            {
                "source": {
                    key: evidence[key] for key in ("file_id", "path", "filename")
                },
                "snapshot": {
                    key: evidence[key] for key in ("size", "mtime_ns", "sha256")
                },
            }
        ]
    }
    plan["selected_files"][0][section][field] = value
    before = deepcopy((plan, checked))
    with pytest.raises(ValueError, match="snapshot differs"):
        bind_selected_snapshots(plan, checked)
    assert (plan, checked) == before


def test_valid_index_does_not_establish_task_provenance(tmp_path, source_read):
    _, original = source_read
    # There is no TaskSnapshot argument: an otherwise valid arbitrary index is
    # preserved, not authenticated as current, in-range, or complete evidence.
    checked = revalidate_documents(str(tmp_path), [(999, original)])
    assert set(checked) == {999}


@pytest.fixture
def c_package_request(tmp_path):
    workspace = tmp_path / "中文父目录 空格" / "课程材料"
    workspace.mkdir(parents=True)
    source = workspace / "课程笔记.txt"
    source.write_text("课程原始材料", encoding="utf-8")
    result = read_document(str(source))
    assert result["success"], result
    document = result["data"]["document"]
    return {
        "task": {
            "task_id": "c-platform-check",
            "user_goal": "准备课程资料，只读规划",
            "workspace": str(workspace),
        },
        "selected_files": [
            {
                "source": {
                    key: document[key] for key in ("file_id", "path", "filename")
                },
                "requirements": [],
                "reason": "受控测试提议，不代表批准",
            }
        ],
        "destination": str(workspace / "待审批课程包"),
        "analysis": dict.fromkeys(
            [
                "extract_requirements",
                "match_requirements",
                "detect_conflicts",
                "compare_versions",
            ]
        ),
    }


def test_same_length_mutation_during_package_read_restored_mtime_rejects(
    c_package_request, monkeypatch
):
    request = c_package_request
    source = Path(request["selected_files"][0]["source"]["path"])
    before = source.stat()
    raw = source.read_bytes()
    replacement = raw.replace("原始".encode(), "变化".encode())
    assert replacement != raw and len(replacement) == len(raw)
    original_read = os.read
    mutated = False

    def changing_read(fd, size):
        nonlocal mutated
        data = original_read(fd, size)
        if data and not mutated:
            # Modify the real file in place. An OS denial is a test failure,
            # not evidence that the planner detected an actual mutation.
            source.write_bytes(replacement)
            os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns))
            assert source.read_bytes() == replacement
            assert source.stat().st_size == before.st_size
            assert source.stat().st_mtime_ns == before.st_mtime_ns
            mutated = True
        return data

    inputs = deepcopy(request)
    monkeypatch.setattr(os, "read", changing_read)
    result = plan_package(**request)
    assert mutated, result
    assert result["success"] is False and result["data"] is None, result
    assert result["tool"] == "plan_package"
    assert result["message"] and "changed" in result["error"].lower(), result
    assert request == inputs
    assert not Path(request["destination"]).exists()
    assert source.read_bytes() == replacement


def test_chinese_parent_directory_readonly_package_plan(c_package_request):
    request = c_package_request
    source = Path(request["selected_files"][0]["source"]["path"])
    before = (source.read_bytes(), source.stat().st_mtime_ns)
    entries = sorted(p.relative_to(source.parent) for p in source.parent.rglob("*"))
    inputs = deepcopy(request)
    result = plan_package(**request)
    assert result["success"] is True and result["error"] is None, result
    plan = result["data"]["plan"]
    selected = plan["selected_files"][0]
    assert selected["source"] == request["selected_files"][0]["source"]
    assert selected["snapshot"]["sha256"] == hashlib.sha256(before[0]).hexdigest()
    assert selected["snapshot"]["size"] == len(before[0])
    assert selected["snapshot"]["mtime_ns"] == before[1]
    assert plan["destination"] == request["destination"]
    assert request == inputs
    assert before == (source.read_bytes(), source.stat().st_mtime_ns)
    assert entries == sorted(
        p.relative_to(source.parent) for p in source.parent.rglob("*")
    )
    assert not Path(request["destination"]).exists()
